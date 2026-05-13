import json
import io
import time
from typing import List
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..auth import get_current_admin
from ..db import get_db
from ..models import LaterItem, Resource, Subject, User
from ..schemas import BatchDeleteRequest, ResourceListOut, ResourceOut, SubjectCreate, SubjectOut, SubjectUpdate
from ..routers.subjects import resource_to_out

import openpyxl

router = APIRouter(prefix="/admin", tags=["admin"])


TEMPLATE_HEADERS = ["标题", "URL", "描述", "详细介绍", "标签(逗号分隔)", "平台", "所属学科ID"]


@router.get("/template")
def download_template(
    current_admin: User = Depends(get_current_admin),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "资源批量导入"
    ws.append(TEMPLATE_HEADERS)

    # Add example row
    ws.append([
        "示例：数据结构课程笔记",
        "https://example.com/ds-notes",
        "一份详尽的数据结构学习笔记",
        "本资源涵盖了线性表、栈、队列、树、图等核心数据结构，配有大量习题和动画演示，适合考研复习和期末备考。",
        "课程,笔记,考研",
        "GitHub",
        "c",
    ])

    # Set column widths
    widths = [30, 40, 30, 45, 25, 15, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=resource_import_template.xlsx"},
    )


@router.post("/import-excel")
async def import_excel(
    file: UploadFile,
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db),
):
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="请上传 .xlsx 文件")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析 Excel 文件")

    rows = list(ws.iter_rows(min_row=2, values_only=True))  # skip header
    if not rows:
        raise HTTPException(status_code=400, detail="Excel 文件中没有数据行")

    # Preload subjects for validation
    subject_ids = {s.id for s in db.query(Subject).all()}
    # Preload existing resource URLs per subject for dedup
    existing_resources = {}
    for r in db.query(Resource).filter(Resource.client_id == "default").all():
        key = (r.subject_id, r.url.strip())
        existing_resources[key] = r.title

    created = 0
    errors = []
    seen_in_batch = {}
    now = int(time.time() * 1000)

    for i, row in enumerate(rows, start=2):  # row 2 onwards
        if not row or not any(cell for cell in row):
            continue  # skip empty rows
        if len(row) < 7:
            errors.append({"row": i, "error": "列数不足，需要7列"})
            continue

        title, url, description, detailed_desc, tags_str, platform, subject_id = row[:7]

        if not title or not url:
            errors.append({"row": i, "error": "标题和URL为必填项"})
            continue
        if not subject_id:
            errors.append({"row": i, "error": "所属学科ID为必填项"})
            continue

        sid = str(subject_id).strip()
        surl = str(url).strip()

        if sid not in subject_ids:
            errors.append({"row": i, "error": f"学科ID '{sid}' 不存在，请先在管理面板「学科管理」中添加"})
            continue

        # Dedup: within batch
        batch_key = (sid, surl)
        if batch_key in seen_in_batch:
            errors.append({"row": i, "error": f"URL 与本次导入第 {seen_in_batch[batch_key]} 行重复，已跳过"})
            continue

        # Dedup: against existing DB
        if batch_key in existing_resources:
            errors.append({"row": i, "error": f"URL 已存在于资源「{existing_resources[batch_key]}」中，已跳过"})
            continue

        seen_in_batch[batch_key] = i

        resource = Resource(
            id=f"res_{uuid4().hex[:12]}",
            subject_id=sid,
            title=str(title).strip(),
            url=surl,
            description=str(description).strip() if description else None,
            detailed_description=str(detailed_desc).strip() if detailed_desc else None,
            platform=str(platform).strip() if platform else None,
            tags=json.dumps([t.strip() for t in str(tags_str).split(",") if t.strip()], ensure_ascii=False) if tags_str else None,
            created_at=now,
            status="approved",
            visibility="public",
            submitter_id=current_admin.id,
            client_id="default",
        )
        db.add(resource)
        created += 1

    db.commit()

    return {"created": created, "errors": errors}


# ── Resource management (admin only) ──────────────────────


@router.get("/resources", response_model=ResourceListOut)
def admin_list_resources(
    subject_id: str = None,
    keyword: str = None,
    page: int = 1,
    page_size: int = 50,
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db),
):
    """List all public resources with optional subject/keyword filters."""
    query = db.query(Resource).filter(Resource.client_id == "default")

    if subject_id:
        query = query.filter(Resource.subject_id == subject_id)

    if keyword:
        like_term = f"%{keyword.strip()}%"
        query = query.filter(
            or_(
                Resource.title.ilike(like_term),
                Resource.url.ilike(like_term),
                Resource.description.ilike(like_term),
            )
        )

    total = query.count()
    offset = max(0, (page - 1)) * page_size
    items = (
        query.order_by(Resource.created_at.desc())
        .offset(offset)
        .limit(page_size)
        .all()
    )

    return ResourceListOut(
        resources=[resource_to_out(item) for item in items],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.post("/resources/batch-delete")
def admin_batch_delete_resources(
    payload: BatchDeleteRequest,
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db),
):
    """Batch delete public resources by IDs. Only deletes client_id='default' resources."""
    if not payload.resource_ids:
        raise HTTPException(status_code=400, detail="请提供要删除的资源ID列表")

    deleted = 0
    errors = []

    for rid in payload.resource_ids:
        item = db.query(Resource).filter(Resource.id == rid).first()
        if not item:
            errors.append({"resource_id": rid, "error": "资源不存在"})
            continue
        if item.client_id != "default":
            errors.append(
                {"resource_id": rid, "error": "只能删除公共资源，无法删除用户个人资源"}
            )
            continue
        db.delete(item)
        deleted += 1

    db.commit()
    return {"deleted": deleted, "errors": errors}


# ── Subject CRUD (admin only) ────────────────────────────

@router.get("/subjects", response_model=List[SubjectOut])
def admin_list_subjects(
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db),
):
    return db.query(Subject).order_by(Subject.id.asc()).all()


@router.post("/subjects", response_model=SubjectOut, status_code=201)
def create_subject(
    payload: SubjectCreate,
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db),
):
    sid = payload.id.strip()
    if not sid or not payload.name.strip():
        raise HTTPException(status_code=400, detail="学科ID和名称为必填项")
    if db.query(Subject).filter(Subject.id == sid).first():
        raise HTTPException(status_code=400, detail="学科ID已存在")
    subject = Subject(id=sid, name=payload.name.strip())
    db.add(subject)
    db.commit()
    db.refresh(subject)
    return subject


@router.post("/subjects/auto", response_model=SubjectOut, status_code=201)
def auto_create_subject(
    payload: SubjectCreate,
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db),
):
    if not payload.name.strip():
        raise HTTPException(status_code=400, detail="学科名称为必填项")
    # Find the next available sub_N ID
    existing = {s.id for s in db.query(Subject).all()}
    n = 1
    while f"sub_{n}" in existing:
        n += 1
    subject = Subject(id=f"sub_{n}", name=payload.name.strip())
    db.add(subject)
    db.commit()
    db.refresh(subject)
    return subject


@router.patch("/subjects/{subject_id}", response_model=SubjectOut)
def update_subject(
    subject_id: str,
    payload: SubjectUpdate,
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db),
):
    subject = db.query(Subject).filter(Subject.id == subject_id).first()
    if not subject:
        raise HTTPException(status_code=404, detail="学科不存在")
    if payload.id is not None:
        new_id = payload.id.strip()
        if not new_id:
            raise HTTPException(status_code=400, detail="学科ID不能为空")
        if new_id != subject_id and db.query(Subject).filter(Subject.id == new_id).first():
            raise HTTPException(status_code=400, detail="学科ID已存在")
        # Update FK references
        db.query(Resource).filter(Resource.subject_id == subject_id).update(
            {Resource.subject_id: new_id}
        )
        db.query(LaterItem).filter(LaterItem.subject_id == subject_id).update(
            {LaterItem.subject_id: new_id}
        )
        subject.id = new_id
    if payload.name is not None:
        if not payload.name.strip():
            raise HTTPException(status_code=400, detail="学科名称不能为空")
        subject.name = payload.name.strip()
    db.commit()
    db.refresh(subject)
    return subject


@router.delete("/subjects/{subject_id}")
def delete_subject(
    subject_id: str,
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db),
):
    subject = db.query(Subject).filter(Subject.id == subject_id).first()
    if not subject:
        raise HTTPException(status_code=404, detail="学科不存在")
    resource_count = db.query(Resource).filter(
        Resource.subject_id == subject_id, Resource.client_id == "default"
    ).count()
    if resource_count > 0:
        raise HTTPException(
            status_code=400,
            detail=f"该学科下有 {resource_count} 个资源，请先删除或迁移资源",
        )
    db.delete(subject)
    db.commit()
    return {"ok": True}
